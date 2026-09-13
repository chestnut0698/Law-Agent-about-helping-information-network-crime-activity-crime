"""报告导出：把服务端存储的 markdown 报告即时转成公文版式 docx，不落盘、不改存储。

版式按 GB/T 9704 常用项：A4；上 37mm、下 35mm、左 28mm、右 26mm；
文题黑体二号居中；节题黑体三号；小节楷体三号；正文仿宋三号、首行缩进两字、固定行距 28 磅。
"""

from __future__ import annotations

import io
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

_INLINE_MARK_RE = re.compile(r"(\*\*|__|\*|_|`)")
_LINK_RE = re.compile(r"\[([^\]]+)\]\([^)]+\)")
_ORDERED_RE = re.compile(r"^\d+[.)]\s+")

_FONT_LATIN = "Times New Roman"
_FONT_TITLE = "黑体"
_FONT_H3 = "楷体_GB2312"
_FONT_BODY = "仿宋_GB2312"
_BLACK = RGBColor(0x00, 0x00, 0x00)
_BODY_PT = 16
_TITLE_PT = 22
_LINE_PT = 28
_INDENT = Pt(32)


def _clean(text: str) -> str:
    text = _LINK_RE.sub(r"\1", text)
    return _INLINE_MARK_RE.sub("", text).strip()


def parse_report_blocks(markdown: str) -> list[tuple[str, str]]:
    """把核验单 markdown 拆成 (kind, text)，导出与页内预览共用同一套切分。"""
    blocks: list[tuple[str, str]] = []
    for raw in (markdown or "").splitlines():
        stripped = raw.strip()
        if not stripped or stripped.startswith("```"):
            continue
        if len(stripped) >= 3 and set(stripped) <= set("-—*_"):
            continue
        if stripped.startswith("### "):
            blocks.append(("h3", _clean(stripped[4:])))
        elif stripped.startswith("## "):
            blocks.append(("h2", _clean(stripped[3:])))
        elif stripped.startswith("# "):
            blocks.append(("h1", _clean(stripped[2:])))
        elif stripped.startswith(("- ", "* ", "• ")):
            blocks.append(("ul", _clean(stripped[2:])))
        elif _ORDERED_RE.match(stripped):
            blocks.append(("ol", _clean(_ORDERED_RE.sub("", stripped))))
        elif stripped.startswith(">"):
            blocks.append(("body", _clean(stripped.lstrip(">").strip())))
        else:
            blocks.append(("body", _clean(stripped)))
    return blocks


def _set_run_font(run, east_asia: str, size_pt: float, *, bold: bool = False) -> None:
    run.bold = bold
    run.italic = False
    run.font.size = Pt(size_pt)
    run.font.color.rgb = _BLACK
    run.font.name = _FONT_LATIN
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), _FONT_LATIN)
    rFonts.set(qn("w:hAnsi"), _FONT_LATIN)
    rFonts.set(qn("w:eastAsia"), east_asia)
    rFonts.set(qn("w:cs"), _FONT_LATIN)


def _set_style_font(style, east_asia: str, size_pt: float, *, bold: bool = False) -> None:
    font = style.font
    font.name = _FONT_LATIN
    font.size = Pt(size_pt)
    font.bold = bold
    font.italic = False
    font.color.rgb = _BLACK
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), _FONT_LATIN)
    rFonts.set(qn("w:hAnsi"), _FONT_LATIN)
    rFonts.set(qn("w:eastAsia"), east_asia)
    rFonts.set(qn("w:cs"), _FONT_LATIN)


def _configure_page(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(37)
    section.bottom_margin = Mm(35)
    section.left_margin = Mm(28)
    section.right_margin = Mm(26)


def _configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    _set_style_font(normal, _FONT_BODY, _BODY_PT, bold=False)
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(_LINE_PT)
    pf.space_after = Pt(0)
    for name in ("List Bullet", "List Number"):
        try:
            _set_style_font(doc.styles[name], _FONT_BODY, _BODY_PT, bold=False)
        except KeyError:
            continue


def _apply_line_metrics(paragraph, *, exact: bool = True) -> None:
    pf = paragraph.paragraph_format
    if exact:
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = Pt(_LINE_PT)
    else:
        pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(0)


def _add_paragraph(doc: Document, text: str, *, kind: str = "body"):
    style = None
    if kind == "ul":
        style = "List Bullet"
    elif kind == "ol":
        style = "List Number"
    p = doc.add_paragraph(text, style=style) if style else doc.add_paragraph(text)
    pf = p.paragraph_format
    if kind == "h1":
        east, size, bold = _FONT_TITLE, _TITLE_PT, True
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.first_line_indent = Pt(0)
        pf.space_before = Pt(0)
        pf.space_after = Pt(12)
        _apply_line_metrics(p, exact=False)
    elif kind == "h2":
        east, size, bold = _FONT_TITLE, _BODY_PT, True
        pf.first_line_indent = Pt(0)
        pf.space_before = Pt(12)
        pf.space_after = Pt(6)
        _apply_line_metrics(p)
    elif kind == "h3":
        east, size, bold = _FONT_H3, _BODY_PT, False
        pf.first_line_indent = Pt(0)
        pf.space_before = Pt(8)
        pf.space_after = Pt(4)
        _apply_line_metrics(p)
    elif kind in {"ul", "ol"}:
        east, size, bold = _FONT_BODY, _BODY_PT, False
        pf.first_line_indent = Pt(0)
        pf.left_indent = _INDENT
        _apply_line_metrics(p)
    else:
        east, size, bold = _FONT_BODY, _BODY_PT, False
        pf.first_line_indent = _INDENT
        pf.left_indent = Pt(0)
        _apply_line_metrics(p)
    if not p.runs:
        run = p.add_run(text)
        _set_run_font(run, east, size, bold=bold)
    else:
        for run in p.runs:
            _set_run_font(run, east, size, bold=bold)
    return p


def markdown_to_docx(markdown: str, title: str = "跨案关联线索核验单") -> bytes:
    doc = Document()
    _configure_page(doc)
    _configure_styles(doc)
    try:
        doc.core_properties.title = title
    except Exception:
        pass

    for kind, text in parse_report_blocks(markdown):
        _add_paragraph(doc, text, kind=kind)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
